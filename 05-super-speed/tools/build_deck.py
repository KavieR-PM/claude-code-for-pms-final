"""Build deck.pptx (Rook Signal decision deck) with speaker notes.

Run from the repo root with the project-local environment:
    scratch/.venv/bin/python 05-super-speed/tools/build_deck.py
Content matches deck.html; keep the two in sync.
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "deck.pptx"

NIGHT, SKY, RED, YELLOW, GREEN = "0D1B4C", "1F4FD8", "E3262F", "FFCC1F", "13A862"
INK, PAPER, PAPER2, MUTED, SOFT = "0B0F1F", "FFFFFF", "EEF3FF", "4D5778", "C9D4FF"
HEAD, BODY = "Arial Black", "Calibri"
W, H = 13.333, 7.5
M = 0.6

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
TOTAL = 13
count = 0


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=2.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, bold)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        parts = para if isinstance(para, list) else [(para, bold)]
        for t, b in parts:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = b; r.font.name = font; r.font.color.rgb = rgb(color)
    return tb


def new_slide(dark, notes, kicker, title=None, title_size=32):
    global count
    count += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid(); bg.fore_color.rgb = rgb(NIGHT if dark else PAPER)
    text(s, M, 0.45, 9, 0.35, kicker.upper(), size=13, bold=True, color=SOFT if dark else RED)
    if title:
        text(s, M, 0.8, W - 2 * M, 0.9, title, size=title_size, font=HEAD, color=YELLOW if dark else INK)
    text(s, W - 1.4, H - 0.45, 0.9, 0.3, f"{count} / {TOTAL}", size=11, color=SOFT if dark else MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def source(slide, words, dark=False):
    text(slide, M, H - 0.5, W - 2.4, 0.3, words, size=11, color=SOFT if dark else MUTED)


def card(slide, x, y, w, h, title, body, fill=PAPER2, fg=INK, title_color=None):
    box(slide, x + 0.07, y + 0.07, w, h, fill=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(slide, x, y, w, h, fill=fill, line=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x + 0.25, y + 0.2, w - 0.5, 0.45, title, size=16, font=HEAD, color=title_color or fg)
    text(slide, x + 0.25, y + 0.72, w - 0.5, h - 0.85, body, size=15, color=fg)


def column_chart(slide, x, y, w, h, cats, vals, hot_from, label_color):
    data = CategoryChartData(); data.categories = cats; data.add_series("Offers per week", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(12)
    ch.category_axis.tick_labels.font.color.rgb = rgb(label_color)
    ch.category_axis.format.line.fill.background()
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(13); dl.font.bold = True; dl.font.color.rgb = rgb(label_color)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(RED if i >= hot_from else SKY)
        pt.format.line.color.rgb = rgb(INK)
    return ch


def table(slide, x, y, w, rows, col_w, size=16):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.5 * len(rows)))
    tbl = shp.table
    for c, cw in enumerate(col_w):
        tbl.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(INK if r == 0 else (PAPER2 if r % 2 else PAPER))
            tf = cell.text_frame; tf.word_wrap = True
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run(); run.text = val
            run.font.size = Pt(13 if r == 0 else size); run.font.bold = r == 0 or c == 0 and False
            run.font.name = BODY; run.font.color.rgb = rgb(PAPER if r == 0 else INK)
    return tbl


# 1. Title
s = new_slide(True, "This is a decision meeting. In about fifteen minutes I want to show you what 4.2 actually did to responders, why it happened, and what I think we should build, so you can make three calls at the end. Everything here comes from the callout history, 25 support tickets, Sofia's four handler interviews and the routing code. One caveat up front: the callout file hasn't been verified by Ravi yet, so treat the exact numbers as strong signals, not final figures.", "Rook Dispatch · for Helen Achebe · 8 Oct 2026")
text(s, M, 1.6, 11, 2.6, ["Bringing quiet", "responders back"], size=60, font=HEAD, color=YELLOW)
text(s, M, 4.4, 10, 0.6, "What 4.2 did, why, and the fix we'd build: Rook Signal.", size=22, color=PAPER)
source(s, "Owners: Dispatch PM · Wen Li & Marcus Oyelaran's team · Sofia Marino", dark=True)

# 2. The one number
s = new_slide(True, "If you remember one thing, it's this. Four of our sixteen responders went from about twelve offers a week each to almost none. Combined, they went from 49 offers a week to 3, and not one of them has come back. These weren't weak responders; before the release they accepted 69 to 87 percent of their offers, and Vesper was one of the best in the file. The chart shows the four combined, week by week. The drop starts the week 4.2 shipped and never recovers.", "The one number")
text(s, M, 1.4, 6.2, 2.0, "4 of 16", size=80, font=HEAD, color=YELLOW)
text(s, M, 3.7, 5.6, 1.6, [[("responders went from ~12 offers a week to ", False), ("0–1", True), (". Combined: ", False), ("49 → 3", True), (". None has recovered.", False)]], size=22, color=PAPER)
column_chart(s, 7.0, 1.3, 5.7, 4.8, ["Before", "10 Aug", "17 Aug", "24 Aug", "31 Aug"], [49, 43, 16, 6, 3], 2, PAPER)
source(s, "Source: callout history CSV: Farlight, Meteor Mite, The Undertow, Vesper (not yet verified by Ravi)", dark=True)

# 3. Vesper
s = new_slide(False, "Here's what that looks like for one person. Vesper took eight of every ten jobs all summer. His handler, Aunt Dot, told Sofia that his phone often buzzes downstairs while he's upstairs. With 90 seconds he made it. With 60 he didn't. He missed six jobs in the release week, each miss moved him down the line, and fewer offers came. By the end of August he was getting one offer a week, and nothing told him why. Dot only noticed because she was cooking for someone who was home all week.", "One responder's August", "Vesper: 14 offers a week → 1")
column_chart(s, M, 1.9, 6.6, 4.4, ["29 Jun", "6 Jul", "13 Jul", "20 Jul", "27 Jul", "3 Aug", "10 Aug", "17 Aug", "24 Aug", "31 Aug"], [14, 13, 15, 14, 13, 14, 12, 5, 2, 1], 6, INK)
text(s, 7.7, 2.4, 5.0, 2.6, "“By the time he’s actually got a thumb on the screen — it’s gone. Gone to somebody else already.”", size=26, bold=True)
text(s, 7.7, 5.2, 5.0, 0.5, "Aunt Dot, Vesper's handler", size=16, color=MUTED)
source(s, "Weeks of 29 Jun – 31 Aug; 4.2 shipped Wed 12 Aug. Source: callout history CSV; Sofia's interview, 3 Sep")

# 4. How it happened
s = new_slide(False, "Three things combined. First, the shorter timer made everyone miss more in the release week: acceptance fell from 78 to 54 percent. Second, the scoring rules. A missed offer counts exactly like a no, a miss costs more than a yes earns, so you need six out of ten just to stand still, and nothing ever restores points. Third, offers go down the line one at a time, so if you're at the back you're only asked when everyone ahead says no. Together, the people who missed the most fell to the back and could never earn their way out.", "How it happened", "A bad week, then a trap with no exit")
cw = 3.55
for i, (t, b, f, fg) in enumerate([
    ("1. Shorter timer", "4.2 cut the answer window from 90 to 60 seconds. Everyone missed more: acceptance fell from 78% to 54%.", RED, PAPER),
    ("2. The points trap", "A miss counts like a “no”. It costs 0.12; a yes earns 0.08. Points never come back on their own.", YELLOW, INK),
    ("3. Back of the line", "Offers go one at a time. At the back, you’re only asked when everyone ahead says no, so there’s no way up.", SKY, PAPER),
]):
    x = M + i * (cw + 0.65)
    card(s, x, 2.2, cw, 3.4, t, b, fill=f, fg=fg)
    if i < 2:
        text(s, x + cw + 0.1, 3.5, 0.45, 0.7, "→", size=36, bold=True, color=RED, align=PP_ALIGN.CENTER)
source(s, "Source: dispatch-routing code (offer.py:30, history.py, config.py); callout history CSV")

# 5. Was it August
s = new_slide(False, "Priya's handover said this was August being quiet. The data doesn't support that for the collapse. The first week of August was the best week of the summer, and the drop came all at once in the release week. A quiet month would hit everyone; instead twelve responders got more work and four got none. The clearest proof is Kip's two responders: same handler, same city, same weeks. Meteor Mite went from eleven offers to one while The Gale went from thirteen to twenty-one. August might explain a small slice of the lower volume afterwards, at most around thirty percent, but none of the collapse.", "Was it just August?", "No: same city, opposite fates")
cw = 3.75
for i, (t, b, f) in enumerate([
    ("No early dip", "Week of 3 Aug was the summer’s best (78%). The drop came all at once, in the release week.", PAPER2),
    ("Not everyone", "12 responders got more work. 4 got almost none. A quiet month doesn’t pick four people.", PAPER2),
    ("Kip’s two", "Same handler, same city: Meteor Mite 11 → 1, The Gale 13 → 21.", YELLOW),
]):
    card(s, M + i * (cw + 0.42), 2.1, cw, 2.9, t, b, fill=f)
text(s, M, 5.45, 12, 0.5, "August explains 0% of the collapse, and at most ~30% of the lower volume after it.", size=18, color=MUTED)
source(s, "Source: callout history CSV; four-analyst root-cause review")

# 6. Why nobody saw it
s = new_slide(False, "So why didn't anyone see this? Three reasons. Our headline metric hid it: acceptance rate looks like it recovered to 73 percent, but partly because the stuck responders stopped being offered anything, so their misses dropped out of the maths. Callouts actually filled are still about nine percent down. Second, a missed offer just disappears from the phone, so about forty percent of tickets are really asking 'is my account broken?' Third, two of the four stuck responders never filed a ticket at all; we only found them in Sofia's design interviews.", "Why nobody saw it", "The dashboard said “recovered”")
for i, (t, b, f, fg) in enumerate([
    ("73% “recovered”", "Acceptance looks almost normal because the stuck four stopped being asked. Callouts filled are still 9% down.", SKY, PAPER),
    ("40% of tickets", "ask “is my account broken?” Missed offers vanish from the phone with no trace.", YELLOW, INK),
    ("2 of 4 silent", "Mite and Vesper never filed a ticket. They surfaced only in design interviews.", RED, PAPER),
]):
    card(s, M + i * (3.75 + 0.42), 2.1, 3.75, 3.2, t, b, fill=f, fg=fg)
source(s, "Source: callout history CSV; 25 support tickets; Sofia's handler interviews")

# 7. Root cause and confidence
s = new_slide(False, "Four analysts looked at this from different angles and debated it. They agree with high confidence that the scoring rules are the lock-in and must be fixed. They agree with medium confidence that the shorter timer was the trigger. What they could not settle is why those four specifically kept missing afterwards. It's either that the timer is too short for how they answer, that their phones weren't getting offers in time, or that they chose not to. The fix I'm proposing works for all three, and the new logging will tell us which it is.", "Root cause, and how sure we are", "What we know, and what's still open")
table(s, M, 2.0, 12.1, [
    ["Finding", "Confidence"],
    ["The scoring rules lock out anyone who keeps missing", "High"],
    ["The 60-second timer triggered the bad week", "Medium"],
    ["The closeness reweight moved work between others; it didn't cause the collapse", "Medium"],
    ["Open: why those four kept missing (timer too short, phones late, or choice)", "Low"],
], [9.6, 2.5], size=18)
source(s, "Source: four-analyst root-cause debate; hypotheses #1–#11 in the analysis file")

# 8. What we'd build
s = new_slide(False, "Here's what we'd build, from the point of view of the people it happens to. A missed offer stops counting against you; a real no still does, which keeps Wen's original intent. Every offer says what, where and how far, with leave-by and arrive-by times. A fifteen-second grace window gives near misses a fair chance without slowing dispatch. Every miss comes back with context and a one-tap reason, which is also how we'll learn whether it's the timer or the phones. Responders can say 'I'm free now, put me first', with limits. And handlers see who's quiet or overloaded, and can act.", "What we'd build", "Rook Signal: six changes")
items = [
    ("Missing isn’t “no”", "Misses cost nothing; a “no” still counts. The four get a fresh start."),
    ("Where and how far", "Drive time first, leave-by and arrive-by, map, skills, a giant Take It button."),
    ("15-second grace", "“Still lit!” First tap wins. Dispatch doesn’t slow down."),
    ("Signal missed", "Context plus a one-tap reason, which tells us timer vs phone."),
    ("Put me first", "“I’m free now.” Once every 4 hours, within a 15-minute drive."),
    ("Handler HQ", "Gone quiet and overloaded flags, with reasons and actions."),
]
cw, chh = 3.75, 2.05
for i, (t, b) in enumerate(items):
    x = M + (i % 3) * (cw + 0.42); y = 1.95 + (i // 3) * (chh + 0.35)
    card(s, x, y, cw, chh, f"{i + 1}. {t}", b)

# 9. Prototype
s = new_slide(True, "This is the clickable prototype; I'll switch to it live. On the left is what Vesper sees when a signal arrives: urgency, the incident, a nine-minute drive in big type, leave by and arrive by, and a huge Take It button you can hit with gloves on. If the sixty seconds run out, the card turns yellow and says 'still lit' for fifteen seconds. The look is an exploratory direction; Sofia will map it to our real design system. The people in it are real cases from the data and interviews; individual offer times are illustrative.", "The prototype", "Click through it: prototype.html")
box(s, M, 1.85, 3.2, 4.85, fill=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, M + 0.15, 2.0, 2.9, 4.55, fill=RED, line=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, M + 0.35, 2.15, 2.5, 0.4, "SIGNAL!", size=18, font=HEAD, color=YELLOW)
text(s, M + 0.35, 2.6, 2.5, 0.5, "Building collapse", size=14, font=HEAD, color=PAPER)
text(s, M + 0.35, 3.1, 2.5, 0.7, "9 MIN DRIVE", size=22, font=HEAD, color=YELLOW)
text(s, M + 0.35, 3.85, 2.5, 0.35, "4.1 km · Harbor district", size=13, color=PAPER)
text(s, M + 0.35, 4.2, 2.5, 0.35, "Leave by 21:24 · arrive 21:33", size=12, bold=True, color=PAPER)
btn = box(s, M + 0.85, 4.75, 1.5, 1.5, fill=PAPER, line=YELLOW, shape=MSO_SHAPE.OVAL, line_w=6)
text(s, M + 0.85, 5.25, 1.5, 0.5, "TAKE IT", size=16, font=HEAD, color=INK, align=PP_ALIGN.CENTER)
cw, chh = 4.0, 2.05
for i, (t, b, f, fg) in enumerate([
    ("Still lit!", "15 more seconds after time runs out. First tap wins.", YELLOW, INK),
    ("It’s yours!", "Instant result, with when to leave.", GREEN, PAPER),
    ("Signal missed", "One tap: what happened? Never changes your offers.", PAPER2, INK),
    ("Handler HQ", "Mite “Gone quiet”, The Gale “Overloaded 1.6×”.", SKY, PAPER),
]):
    card(s, 4.4 + (i % 2) * (cw + 0.35), 1.95 + (i // 2) * (chh + 0.35), cw, chh, t, b, fill=f, fg=fg)
source(s, "Weekly numbers from the callout history; offer details illustrative. Visual style is exploratory.", dark=True)

# 10. Decisions
s = new_slide(False, "These are the three calls I need from you. First, a missed offer should cost nothing, while a real no still counts. If filling a callout slows by more than ten seconds, we revisit. Second, reset the four stuck responders now rather than waiting for more data; they've lost work every week since the twelfth of August. We'd ask Dot and Kip to confirm their phones work first. Third, keep 'put me first' tight: fifteen-minute drive, matching skills, thirty minutes, once every four hours. That protects urgent callouts.", "Decisions", "Three calls I need from you")
table(s, M, 1.95, 12.1, [
    ["Decision", "My recommendation"],
    ["What a missed offer costs", "Nothing; a “no” still counts. Revisit if filling slows by 10+ seconds."],
    ["Reset the four now, or wait for data", "Now, once Dot and Kip confirm the phones work."],
    ["“Put me first” limits", "Tight: 15-min drive, matching skills, 30 min, once every 4 hours."],
], [4.6, 7.5], size=18)
text(s, M, 4.6, 12, 0.9, "Also: room in the next release, Sofia’s time, a yes to writing “how routing decides”, and 10 minutes on Availability Confidence.", size=17, color=MUTED)

# 11. Success
s = new_slide(False, "Here's how we'll know it worked. The four are back to at least half their usual offers within three weeks. Callouts filled per week return to about 132, and we report that next to acceptance rate from now on, because acceptance rate alone hid this. 'Is my account broken' tickets fall by half. And two guardrails: filling a callout shouldn't get more than ten seconds slower, and first-in-line offers should be accepted at close to the normal rate.", "Success", "How we'll know it worked")
cw, chh = 5.85, 1.9
for i, (t, b, f, fg) in enumerate([
    ("4 of 4 back", "Stuck responders at ≥50% of their usual offers within 3 weeks.", GREEN, PAPER),
    ("~132 filled / week", "Callouts filled back to normal (now ~120), reported next to acceptance rate.", SKY, PAPER),
    ("Half the tickets", "“Is my account broken?” tickets fall by half.", YELLOW, INK),
    ("Guardrails", "Filling a callout slows by ≤10 s. First-in-line accept rate stays near normal.", PAPER2, INK),
]):
    card(s, M + (i % 2) * (cw + 0.4), 1.95 + (i // 2) * (chh + 0.35), cw, chh, t, b, fill=f, fg=fg)

# 12. Rollout
s = new_slide(False, "We'd ship in five phases, each behind an on/off switch and piloted in Kip's city first, with Wen owning a two-week watch after each. Logging comes first because it's invisible, small, and settles the timer-versus-phone question. Then the scoring fix and the reset, which stops the harm. One risk to check before anything ships: scores are held in memory in the code we reviewed. If installing a release resets everyone to the middle, the fix itself could restart the spiral. Wen confirms that before we ship.", "Rollout", "Five phases, one risk to check first")
steps = [("1 · Logging", "Every offer: sent, delivered, seen, answered, reason."), ("2 · Fair scoring", "Misses cost nothing; reset the four."), ("3 · The signal", "Context, leave-by, grace window."), ("4 · Missed + controls", "Reasons, mission log, put me first."), ("5 · Handler HQ", "Flags, reasons, actions, alerts.")]
sw = 2.25
for i, (t, b) in enumerate(steps):
    x = M + i * (sw + 0.21)
    box(s, x + 0.06, 1.96, sw, 2.0, fill=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(s, x, 1.9, sw, 2.0, fill=PAPER2, line=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.18, 2.05, sw - 0.36, 0.75, t, size=14, font=HEAD, color=RED)
    text(s, x + 0.18, 2.85, sw - 0.36, 1.0, b, size=14)
card(s, M, 4.35, 12.1, 1.6, "Check first", "If installing a release resets every score to the middle, the fix could restart the spiral. Wen confirms before we ship.", fill=RED, fg=PAPER)
source(s, "Each phase behind an on/off switch, piloted in Kip's city, with a two-week watch owned by Wen.")

# 13. Next steps
s = new_slide(True, "To close: I'm asking for the three decisions, room in the next release for logging and the scoring fix, and Sofia's time. In parallel, I'm sending Wen one combined data request, the offer log, real scores and the push-notification change, and asking Ravi to confirm what the callout file measures and to send September's numbers. The detail is all in the one-pager, the PRD and the engineering handoff. Happy to take questions.", "Next steps", "What happens next")
for i, (t, b, f, fg) in enumerate([
    ("Today", "Your three decisions, release room, Sofia’s time.", YELLOW, INK),
    ("This week", "Wen: offer log, real scores, the push change. Ravi: CSV definitions, September data.", PAPER2, INK),
    ("Next release", "Phase 1 logging and phase 2 fair scoring, piloted in Kip’s city.", GREEN, PAPER),
]):
    card(s, M + i * (3.75 + 0.42), 2.1, 3.75, 3.0, t, b, fill=f, fg=fg)
source(s, "Detail: brief.md (one-pager) · prd.md · engineering-handoff.md · prototype.html", dark=True)

assert count == TOTAL, count
prs.save(OUT)
print("saved", OUT)
